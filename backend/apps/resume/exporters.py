"""Resume exporters: PDF (reportlab) and DOCX (python-docx)."""
import io

from core.utils import slugify

from .models import ResumeVersion


def export_pdf(resume: ResumeVersion) -> tuple[bytes, str, str]:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title=f"{resume.full_name} Resume")
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Title"], fontSize=20, textColor=colors.HexColor("#0f3460"))
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=13, textColor=colors.HexColor("#0f3460"))
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontSize=10, leading=14)

    story = [Paragraph(resume.full_name, h1), Paragraph(resume.title, h2), Spacer(1, 8)]
    if resume.summary:
        story += [Paragraph(resume.summary, body), Spacer(1, 8)]
    for section, items in (("Experience", resume.experience), ("Projects", resume.projects)):
        if not items:
            continue
        story.append(Paragraph(section, h2))
        for item in items:
            header = item.get("title") or item.get("name") or ""
            company = item.get("company", "")
            story.append(Paragraph(f"<b>{header}</b> — {company}", body))
            for b in item.get("bullets", []):
                story.append(Paragraph(f"• {b}", body))
        story.append(Spacer(1, 6))
    if resume.skills:
        story.append(Paragraph("Skills", h2))
        story.append(Paragraph(", ".join(resume.skills), body))
    doc.build(story)
    return buffer.getvalue(), "application/pdf", f"{slugify(resume.full_name)}-resume-v{resume.version_number}.pdf"


def export_docx(resume: ResumeVersion) -> tuple[bytes, str, str]:
    from docx import Document

    document = Document()
    document.add_heading(resume.full_name, level=0)
    document.add_heading(resume.title, level=2)
    if resume.summary:
        document.add_paragraph(resume.summary)
    for section, items in (("Experience", resume.experience), ("Projects", resume.projects)):
        if not items:
            continue
        document.add_heading(section, level=1)
        for item in items:
            document.add_paragraph(f"{item.get('title') or item.get('name', '')} — {item.get('company', '')}")
            for b in item.get("bullets", []):
                document.add_paragraph(b, style="List Bullet")
    if resume.skills:
        document.add_heading("Skills", level=1)
        document.add_paragraph(", ".join(resume.skills))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"{slugify(resume.full_name)}-resume-v{resume.version_number}.docx"
